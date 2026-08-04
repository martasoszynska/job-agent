"""
Orchestrates one full run: fetch -> dedupe against state -> score new ones ->
write output -> update state.

Run manually with: python agent/main.py
"""

import json
from pathlib import Path
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from sources import fetch_all
from score import score_listing

HERE = Path(__file__).parent
PROFILE_PATH = HERE / "profile.json"
STATE_PATH = HERE / "state.json"
OUTPUT_PATH = HERE.parent / "This_Weeks_Picks.xlsx"


def load_json(path):
    with open(path) as f:
        return json.load(f)


def save_json(path, data):
    with open(path, "w") as f:
        json.dump(data, f, indent=2)


def dedupe(listings, state):
    seen_urls = {item["url"] for item in state["seen_links"]}
    new_listings = [l for l in listings if l["url"] not in seen_urls]
    return new_listings


def write_excel(scored_listings):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "This Week's New Picks"
    headers = ["Title", "Organization/Source", "URL", "Match Score", "Notes", "Legal Work Flag"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    for listing, score in scored_listings:
        ws.append([
            listing.get("title"),
            listing.get("source"),
            listing.get("url"),
            score.get("match_score"),
            score.get("notes"),
            score.get("legal_work_flag"),
        ])
    for col, width in zip("ABCDEF", [30, 20, 45, 12, 55, 30]):
        ws.column_dimensions[col].width = width
    wb.save(OUTPUT_PATH)


def main():
    profile = load_json(PROFILE_PATH)
    state = load_json(STATE_PATH)

    print("Fetching listings from all sources...")
    all_listings = fetch_all(profile)
    print(f"Found {len(all_listings)} total listings across sources.")

    new_listings = dedupe(all_listings, state)
    print(f"{len(new_listings)} are new since last run.")

    scored = []
    for listing in new_listings:
        score = score_listing(listing, profile)
        scored.append((listing, score))
        state["seen_links"].append({
            "url": listing["url"],
            "title": listing["title"],
            "organization": listing.get("source", ""),
            "date_first_seen": str(date.today()),
        })

    write_excel(scored)
    save_json(STATE_PATH, state)
    print(f"Wrote {len(scored)} new scored listings to {OUTPUT_PATH}")
    print("Updated state.json - remember this needs to be committed back to the repo.")


if __name__ == "__main__":
    main()
